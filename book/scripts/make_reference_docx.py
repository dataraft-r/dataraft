from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path

root = Path(__file__).resolve().parents[1]
out = root / "reference.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(0.78)
sec.bottom_margin = Inches(0.78)
sec.left_margin = Inches(0.86)
sec.right_margin = Inches(0.74)

styles = doc.styles

normal = styles["Normal"]
normal.font.name = "Liberation Serif"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

for name, size, before, after in [
    ("Title", 28, 0, 18),
    ("Subtitle", 16, 0, 20),
    ("Heading 1", 20, 18, 8),
    ("Heading 2", 14, 14, 5),
    ("Heading 3", 11.5, 10, 4),
]:
    st = styles[name]
    st.font.name = "Liberation Sans"
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0x18, 0x31, 0x53)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
styles["Subtitle"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

for stylename in ["Source Code", "Verbatim Char", "Body Text"]:
    if stylename in styles:
        st = styles[stylename]
        if stylename == "Body Text":
            st.font.name = "Liberation Serif"
            st.font.size = Pt(10.5)
        else:
            st.font.name = "DejaVu Sans Mono"
            st.font.size = Pt(8.3)

# Footer with book title and a PAGE field.
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("Data Platforms with R   |   ")
run.font.name = "Liberation Sans"
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(0x66, 0x6F, 0x7A)

fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
# fldSimple works without an explicit child; Word/LibreOffice resolves it.
footer._p.append(fld)

# A thin divider above the footer.
pPr = footer._p.get_or_add_pPr()
pBdr = OxmlElement("w:pBdr")
top = OxmlElement("w:top")
top.set(qn("w:val"), "single")
top.set(qn("w:sz"), "4")
top.set(qn("w:space"), "4")
top.set(qn("w:color"), "D8DDE3")
pBdr.append(top)
pPr.append(pBdr)

# Keep reference doc minimal. Pandoc imports styles and section properties.
doc.add_paragraph("Reference document for Pandoc output.")
doc.save(out)
print(out)
